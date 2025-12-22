# -*- coding: utf-8 -*-
import logging
import os
import json
import numpy as np
import datetime
from sklearn.manifold import TSNE
from sqlalchemy import text
from dateutil import parser
from openpyxl import Workbook, load_workbook

from seaqa_events.app.config import configs, TEMP_EXPORT_VIEW_DIR
from seaqa_events.db import init_db_session_class
from seaqa_events.app.log import setup_logger
from seaqa_events.utils.connection_utils import list_connection_records_with_columns
from seaqa_events.utils.utils import uuid_str_to_32_chars
from seaqa_events.utils.constant import MAX_EMBEDDING_ANALYSIS_RECORDS, ConnectionType
from seaqa_events.utils.seadb_api import SeaDBAPI
from seaqa_events.utils.sql_view import view_data_2_sql, filter_display_columns

task_logger = setup_logger('seaqa_events_task', propagate=False)

def perform_embedding_analysis(project_uuid, connection_id, connection_type, username):    
    try:
        seadb_api = SeaDBAPI(username)
        column_names = ['_pk', 'title', 'ai_summary', 'ai_summary_vector']
        if connection_type == ConnectionType.GITHUB_ISSUE.value:
            column_names.extend(['url', 'state'])
        records = list_connection_records_with_columns(seadb_api, project_uuid, connection_id, connection_type, column_names,limit=MAX_EMBEDDING_ANALYSIS_RECORDS)
        
        if not records:
            task_logger.warning("No records found for embedding analysis")
            return {'records': []}
            
        valid_records = []
        vectors = []
        
        for record in records:
            ai_summary = record.get('ai_summary')
            ai_summary_vector = record.get('ai_summary_vector')
            if ai_summary and ai_summary_vector:
                record.pop('ai_summary_vector')
                valid_records.append(record)
                vectors.append(ai_summary_vector)

        if not vectors:
            task_logger.warning("No records with embeddings found")
            return {'records': []}
                
        embeddings_2d = generate_embeddings_2d_with_tsne(vectors)
        for i, record in enumerate(valid_records):
            if i < len(embeddings_2d):
                record['x'] = embeddings_2d[i][0]
                record['y'] = embeddings_2d[i][1]
        return {
            'records': valid_records,
            'total_processed': len(valid_records)
        }
    
    except Exception as e:
        task_logger.error(f"Error in embedding analysis: {str(e)}")
        raise


def generate_embeddings_2d_with_tsne(vectors):
    if not vectors:
        return []
    
    embeddings_array = np.array(vectors)

    perplexity = min(30, len(embeddings_array) - 1)
    if perplexity < 1:
        perplexity = 5
    
    tsne = TSNE(n_components=2, random_state=42, perplexity=perplexity)
    embeddings_2d = tsne.fit_transform(embeddings_array)
    
    result = []
    for i in range(len(embeddings_2d)):
        coords = [float(embeddings_2d[i][0]), float(embeddings_2d[i][1])]
        result.append(coords)
    
    return result


def convert_kb_view_to_excel(project_uuid, view_id, username):
    try:
        target_dir = os.path.join(TEMP_EXPORT_VIEW_DIR, str(project_uuid))
        os.makedirs(target_dir, exist_ok=True)

        # Fetch project name and view name
        session_class = init_db_session_class(configs)
        with session_class() as session:
            # Get project name
            project_res = session.execute(text("SELECT name FROM projects WHERE uuid=:p"), {'p': uuid_str_to_32_chars(project_uuid)}).fetchone()
            project_name = project_res[0] if project_res else 'project'

            # Get view name
            view_res = session.execute(text("SELECT details FROM knowledge_base_views WHERE project_uuid=:p"), {'p': uuid_str_to_32_chars(project_uuid)}).fetchone()

        details = json.loads(view_res[0]) if view_res else {'views': []}
        view = next((v for v in details.get('views', []) if v.get('_id') == view_id), {'name': 'view'})
        view_name = view.get('name') or 'view'

        # Fetch columns metadata from SeaDB
        seadb_api = SeaDBAPI(username)
        metadata = seadb_api.get_base_metadata(project_uuid)

        tables_metadata = metadata.get('tables') or []
        kb_table_name = 'knowledge_base'
        table_meta = next((t for t in tables_metadata if t.get('name') == kb_table_name), {})
        columns = table_meta.get('columns', [])
        display_columns = filter_display_columns(columns) or columns

        # Build SQL and query rows from SeaDB
        sql = view_data_2_sql(kb_table_name, display_columns, view, username, 0, 1000, include_deleted=True)
        res = seadb_api.query_rows(project_uuid, sql, convert_keys=False)
        rows = res.get('results', [])

        # Write XLSX
        COLUMN_NAME_MAP = {
            'title': 'Title',
            'content': 'Content',
            'creator': 'Creator',
            'created_time': 'Created time',
            'last_modifier': 'Last modifier',
            'modified_time': 'Last modified time',
        }

        excel_name = f'{project_name}_knowledge_base_{view_name}.xlsx'
        target_path = os.path.join(target_dir, excel_name)
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(view_name or 'view')
        
        header = [COLUMN_NAME_MAP.get(c.get('name'), c.get('name')) for c in (display_columns or [])]
        header_names = [c.get('name') for c in (display_columns or [])]
        header_keys = [c.get('key') for c in (display_columns or [])]

        if header:
            ws.append(header)
        for row in (rows or []):
            row_data = []
            for col_name, col_key in zip(header_names, header_keys):
                val = row.get(col_key)
                if col_name in ('created_time', 'modified_time') and val:
                    try:
                        dt = parser.parse(str(val))
                        val = dt.strftime('%Y-%m-%d %H:%M:%S')
                    except:
                        pass
                row_data.append(val)
            ws.append(row_data)
        wb.save(target_path)
        return {}
    except Exception as e:
        task_logger.error(f'KB export error: {e}')
        raise


def preview_import_kb_from_excel(project_uuid, file_name, limit=20):
    try:
        temp_dir = os.path.join(TEMP_EXPORT_VIEW_DIR, str(project_uuid))
        file_path = os.path.join(temp_dir, file_name)
        wb = load_workbook(file_path, read_only=True)
        ws = wb.active

        headers = []
        rows_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            raise Exception("Empty file")

        COLUMN_NAME_MAP = {
            'title': 'Title',
            'content': 'Content',
        }
        header_to_key = {v.lower(): k for k, v in COLUMN_NAME_MAP.items()}

        column_indexes = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            h = str(header).strip().lower()
            key = header_to_key.get(h)
            if key:
                column_indexes[key] = idx

        if 'title' not in column_indexes or 'content' not in column_indexes:
            return {'file_name': file_name, 'preview_rows': [], 'total_rows': 0}

        preview_rows = []
        total_rows = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            row_obj = {}
            for key, idx in column_indexes.items():
                if idx >= len(row):
                    continue
                value = row[idx]
                if value is None:
                    continue
                row_obj[key] = value
            if 'title' in row_obj and 'content' in row_obj:
                total_rows += 1
                if len(preview_rows) < limit:
                    preview_rows.append({'title': row_obj.get('title'), 'content': row_obj.get('content')})
        return {'file_name': file_name, 'preview_rows': preview_rows, 'total_rows': total_rows}
    except Exception as e:
        task_logger.error(f'Import KB preview error: {e}')
        raise


def import_kb_from_excel(project_uuid, username, file_name):
    try:
        temp_dir = os.path.join(TEMP_EXPORT_VIEW_DIR, str(project_uuid))
        file_path = os.path.join(temp_dir, file_name)
        wb = load_workbook(file_path, read_only=True)
        ws = wb.active

        headers = []
        rows_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            raise Exception("Empty file")

        COLUMN_NAME_MAP = {
            'title': 'Title',
            'content': 'Content',
        }
        header_to_key = {v.lower(): k for k, v in COLUMN_NAME_MAP.items()}

        column_indexes = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            h = str(header).strip().lower()
            key = header_to_key.get(h)
            if key:
                column_indexes[key] = idx

        if 'title' not in column_indexes or 'content' not in column_indexes:
            return {}

        rows_to_insert = []
        seadb_api = SeaDBAPI(username)
        kb_table_name = 'knowledge_base'
        now = datetime.datetime.now(datetime.UTC).isoformat()

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            row_obj = {}
            for key, idx in column_indexes.items():
                if idx >= len(row):
                    continue
                value = row[idx]
                if value is None:
                    continue
                row_obj[key] = value
            if 'title' in row_obj and 'content' in row_obj:
                row_obj['creator'] = username
                row_obj['last_modifier'] = username
                row_obj['created_time'] = now
                row_obj['modified_time'] = now
                row_obj['deleted'] = False
                rows_to_insert.append(row_obj)
            if len(rows_to_insert) >= 1000:
                seadb_api.insert_rows(project_uuid, kb_table_name, rows_to_insert)
                rows_to_insert = []
        if rows_to_insert:
            seadb_api.insert_rows(project_uuid, kb_table_name, rows_to_insert)
    except Exception as e:
        task_logger.error(f'Import KB error: {e}')
        raise
    finally:
        if os.path.exists(file_path):
            os.remove(file_path)
    return {}
